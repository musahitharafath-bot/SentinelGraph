"""
reports/report_generator.py
-------------------------------
Turns the analysis + LLM layers into an actual downloadable security report
(.docx), rather than just narrative text in an API response. This is the
Level 3 gap: the LLM narratives already existed (llm/rag.py), but nothing
rendered them into a document artifact someone could hand to a stakeholder.

Pulls from:
  - analysis/critical_assets.py  (risk + centrality ranked assets)
  - analysis/risk_score.py       (composite risk scores)
  - analysis/attack_paths.py     (internet -> critical routes)
  - llm/rag.py                   (executive summary + path narratives,
                                   skipped gracefully if GROQ_API_KEY isn't set)

Usage:
    python -m reports.report_generator --output reports/output/security_report.docx
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from analysis.critical_assets import rank_critical_assets
from analysis.risk_score import compute_risk_scores
from analysis.attack_paths import summarize_attack_surface

logger = logging.getLogger(__name__)

REPORT_DIR = Path(__file__).resolve().parent / "output"
REPORT_DIR.mkdir(exist_ok=True)

SEVERITY_COLORS = {
    "critical": RGBColor(0xC0, 0x00, 0x00),
    "high": RGBColor(0xE3, 0x6C, 0x09),
    "medium": RGBColor(0xBF, 0x9B, 0x00),
    "low": RGBColor(0x2E, 0x7D, 0x32),
}


def _add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(text, level=level)


def _add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    return table


def _try_llm_section(fn, *args, **kwargs) -> Optional[str]:
    """Runs an llm.rag function, degrading gracefully (returns None) if
    GROQ_API_KEY isn't configured -- the rest of the report still generates
    from graph data alone."""
    try:
        result = fn(*args, **kwargs)
        return result.get("narrative")
    except EnvironmentError as e:
        logger.warning("Skipping LLM section (%s): %s", fn.__name__, e)
        return None
    except Exception as e:
        logger.warning("LLM section %s failed: %s", fn.__name__, e)
        return None


def generate_security_report(
    output_path: Optional[Path] = None,
    top_n: int = 10,
    include_llm_narrative: bool = True,
) -> Path:
    """Builds a full security report as a .docx file and returns its path."""
    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = REPORT_DIR / f"security_report_{stamp}.docx"
    output_path = Path(output_path)

    doc = Document()

    # ---- Title page ----
    title = doc.add_heading("Cyber Attack Graph — Security Risk Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # ---- Executive summary (LLM, optional) ----
    _add_heading(doc, "Executive Summary")
    critical_assets = rank_critical_assets(top_n=top_n)
    narrative = None
    if include_llm_narrative:
        from llm.rag import executive_risk_summary

        narrative = _try_llm_section(executive_risk_summary, top_n=top_n)

    if narrative:
        doc.add_paragraph(narrative)
    else:
        doc.add_paragraph(
            "LLM narrative unavailable (GROQ_API_KEY not configured or request failed) -- "
            "showing structured findings only."
        )
        if critical_assets:
            top = critical_assets[0]
            doc.add_paragraph(
                f"The highest-priority asset is {top['asset_id']} "
                f"(composite score {top['critical_asset_score']}), driven by "
                f"{top['vuln_count']} known vulnerabilities and its structural "
                f"position in the network."
            )

    # ---- Top critical assets ----
    _add_heading(doc, "Top Critical Assets")
    doc.add_paragraph(
        "Ranked by a blend of vulnerability/exposure risk and graph centrality "
        "(PageRank + betweenness) -- i.e. both how dangerous an asset's known "
        "vulnerabilities are, and how structurally important it is to the network."
    )
    if critical_assets:
        _add_table(
            doc,
            ["Asset", "Criticality", "Type", "Score", "Risk", "PageRank", "Betweenness", "Vulns"],
            [
                [
                    a["asset_id"], a["criticality_label"], a["asset_type"],
                    a["critical_asset_score"], a["risk_score"],
                    a["pagerank_score"], a["betweenness_score"], a["vuln_count"],
                ]
                for a in critical_assets
            ],
        )
    else:
        doc.add_paragraph("No asset data found -- has the graph been ingested?")

    # ---- Attack surface ----
    _add_heading(doc, "Attack Surface")
    surface = summarize_attack_surface()
    doc.add_paragraph(
        f"{surface.get('num_attack_routes', 0)} distinct route(s) exist from an "
        f"internet-exposed asset to a critical asset."
    )
    exposed = surface.get("exposed_critical_assets", [])
    if exposed:
        _add_heading(doc, "Exposed Critical/High Assets", level=2)
        _add_table(
            doc,
            ["Asset", "Hostname", "Criticality", "Vuln Count", "Max CVSS"],
            [
                [e["asset_id"], e["hostname"], e["criticality"], e["vuln_count"], e["max_cvss"]]
                for e in exposed
            ],
        )

    paths = surface.get("internet_to_critical_paths", [])
    if paths:
        _add_heading(doc, "Internet → Critical Attack Routes", level=2)
        _add_table(
            doc,
            ["From", "To", "Hops", "Path"],
            [
                [p["from_hostname"], p["to_hostname"], p["hops"], " → ".join(p["path"])]
                for p in paths
            ],
        )

        if include_llm_narrative:
            narrative = _try_llm_section(
                __import__("llm.rag", fromlist=["explain_top_attack_routes"]).explain_top_attack_routes,
                limit=3,
            )
            if narrative:
                _add_heading(doc, "Attack Route Risk Narrative", level=2)
                doc.add_paragraph(narrative)

    # ---- Full risk score table ----
    _add_heading(doc, "Full Asset Risk Scores")
    all_scores = compute_risk_scores()
    if all_scores:
        _add_table(
            doc,
            ["Asset", "Criticality", "Type", "Risk Score", "Vuln Count", "Max CVSS"],
            [
                [s["asset_id"], s["criticality"], s["asset_type"], s["risk_score"], s["vuln_count"], s["max_cvss"]]
                for s in all_scores
            ],
        )

    # ---- Footer ----
    doc.add_page_break()
    footer = doc.add_paragraph(
        "This report was generated automatically from the cyber-attack-graph "
        "platform's Neo4j graph, risk-scoring, and centrality analysis, with "
        "narrative sections produced by an LLM where available. Verify "
        "findings before acting on them."
    )
    footer.runs[0].italic = True

    doc.save(str(output_path))
    logger.info("Report saved to %s", output_path)
    return output_path


if __name__ == "__main__":
    import argparse

    logging.basicConfig(level=logging.INFO)
    parser = argparse.ArgumentParser(description="Generate a security risk report (.docx)")
    parser.add_argument("--output", type=str, default=None)
    parser.add_argument("--top-n", type=int, default=10)
    parser.add_argument("--no-llm", action="store_true", help="Skip LLM narrative sections")
    args = parser.parse_args()

    path = generate_security_report(
        output_path=Path(args.output) if args.output else None,
        top_n=args.top_n,
        include_llm_narrative=not args.no_llm,
    )
    print(f"Report generated: {path}")
